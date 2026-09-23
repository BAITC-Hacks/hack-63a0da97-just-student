from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, BadZipFile

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB


class AttachmentError(Exception):
    pass


def extract_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(data))
    except Exception as exc:
        raise AttachmentError("Не удалось открыть PDF") from exc

    text_parts = []

    for index, page in enumerate(reader.pages[:50], start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""

        if text.strip():
            text_parts.append(
                f"--- Страница {index} ---\n{text}"
            )

    result = "\n\n".join(text_parts).strip()

    if not result:
        raise AttachmentError(
            "В PDF не удалось найти текст. "
            "Возможно, это скан."
        )

    return result


def extract_docx(data: bytes) -> str:
    try:
        document = Document(BytesIO(data))
    except Exception as exc:
        raise AttachmentError("Не удалось открыть DOCX") from exc

    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    # Таблицы тоже важны, потому что спецификации
    # часто находятся именно там.
    for table in document.tables:
        for row in table.rows:
            values = []

            for cell in row.cells:
                value = cell.text.strip()

                if value:
                    values.append(value)

            if values:
                parts.append(" | ".join(values))

    result = "\n".join(parts).strip()

    if not result:
        raise AttachmentError(
            "DOCX не содержит читаемого текста."
        )

    return result


def extract_xlsx(data: bytes) -> str:
    try:
        workbook = load_workbook(
            BytesIO(data),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise AttachmentError("Не удалось открыть XLSX") from exc

    parts = []

    try:
        # Ограничиваем количество листов.
        if len(workbook.worksheets) > 5:
            raise AttachmentError("Разделите Excel: максимум 5 листов на один файл.")
        for sheet in workbook.worksheets[:5]:
            parts.append(f"--- Лист: {sheet.title} ---")

            row_count = 0

            for row in sheet.iter_rows(values_only=True):
                row_count += 1

                if row_count > 1000:
                    raise AttachmentError("Разделите Excel: максимум 1000 строк на лист.")

                if any(v is not None for v in row[50:]):
                    raise AttachmentError("Разделите Excel: максимум 50 столбцов на лист.")

                values = []

                for value in row[:50]:
                    if value is None:
                        values.append("")
                    else:
                        values.append(str(value).strip())

                # Убираем пустые строки
                if any(values):
                    parts.append(" | ".join(values))
    finally:
        workbook.close()

    result = "\n".join(parts).strip()

    if not result:
        raise AttachmentError(
            "Excel-файл не содержит данных."
        )

    return result


def extract_attachment(
    filename: str,
    data: bytes,
) -> dict:
    if not filename:
        raise AttachmentError("У файла отсутствует имя.")

    if not data:
        raise AttachmentError("Файл пуст.")

    if len(data) > MAX_FILE_SIZE:
        raise AttachmentError(
            "Файл слишком большой. Максимум 10 MB."
        )

    extension = Path(filename).suffix.lower()
    if extension in (".docx", ".xlsx"):
        try:
            with ZipFile(BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                    raise AttachmentError("Слишком большой распакованный документ (максимум 50 MB).")
        except BadZipFile as exc:
            raise AttachmentError("Повреждённый DOCX/XLSX.") from exc

    if extension == ".pdf":
        text = extract_pdf(data)

    elif extension == ".docx":
        text = extract_docx(data)

    elif extension == ".xlsx":
        text = extract_xlsx(data)

    elif extension == ".xls":
        try:
            import xlrd
            workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
            parts = []
            try:
                for sheet in workbook.sheets():
                    if sheet.nrows > 1000 or sheet.ncols > 50:
                        raise AttachmentError("Разделите XLS: максимум 1000 строк и 50 столбцов на лист.")
                    parts.extend(" | ".join(str(v) for v in sheet.row_values(i)) for i in range(sheet.nrows))
            finally:
                workbook.release_resources()
            text = "\n".join(parts)
        except AttachmentError:
            raise
        except Exception as exc:
            raise AttachmentError("Не удалось открыть XLS.") from exc

    else:
        raise AttachmentError(
            "Формат пока не поддерживается. "
            "Используй PDF, DOCX, XLS или XLSX."
        )

    # Защита от слишком большого текста
    max_chars = 50_000

    truncated = "[Остальные строки пропущены]" in text
    if extension == ".pdf":
        truncated = len(PdfReader(BytesIO(data)).pages) > 50

    if len(text) > max_chars:
        raise AttachmentError("Разделите документ: максимум 50 000 символов на файл.")

    return {
        "filename": filename,
        "extension": extension,
        "text": text,
        "characters": len(text),
        "truncated": truncated,
    }
