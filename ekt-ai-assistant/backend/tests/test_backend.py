import unittest
from io import BytesIO

from docx import Document
from openpyxl import Workbook

from app.models.specification import SpecificationItem
from app.services.attachment_service import (
    AttachmentError,
    MAX_FILE_SIZE,
    extract_attachment,
)
from app.services.catalog_service import CatalogService


def make_pdf(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 50 750 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n"
        + stream
        + b"\nendstream",
    ]
    result = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(result))
        result.extend(f"{index} 0 obj\n".encode("ascii"))
        result.extend(obj)
        result.extend(b"\nendobj\n")
    xref = len(result)
    result.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    result.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        result.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    result.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref}\n%%EOF\n"
        ).encode("ascii")
    )
    return bytes(result)


class AttachmentTests(unittest.TestCase):
    def test_xlsx_rows_are_extracted(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Артикул", "Наименование", "Количество"])
        sheet.append(["027228", "Legrand DRX250 160A 18kA", 5])
        output = BytesIO()
        workbook.save(output)
        workbook.close()

        result = extract_attachment("test.xlsx", output.getvalue())

        self.assertIn("027228 | Legrand DRX250 160A 18kA | 5", result["text"])
        self.assertFalse(result["truncated"])

    def test_docx_table_is_extracted(self):
        document = Document()
        table = document.add_table(rows=1, cols=3)
        cells = table.rows[0].cells
        cells[0].text = "RM35BA10"
        cells[1].text = "Реле контроля насоса Schneider"
        cells[2].text = "2"
        output = BytesIO()
        document.save(output)

        result = extract_attachment("test.docx", output.getvalue())

        self.assertIn("RM35BA10 | Реле контроля насоса Schneider | 2", result["text"])

    def test_text_pdf_is_extracted(self):
        result = extract_attachment(
            "test.pdf",
            make_pdf("027230 Legrand DRX250 160A 25kA 3 pcs"),
        )

        self.assertIn("027230 Legrand DRX250 160A 25kA 3 pcs", result["text"])

    def test_unsupported_and_large_files_are_rejected(self):
        with self.assertRaisesRegex(AttachmentError, "не поддерживается"):
            extract_attachment("test.txt", b"data")
        with self.assertRaisesRegex(AttachmentError, "слишком большой"):
            extract_attachment("test.pdf", b"x" * (MAX_FILE_SIZE + 1))


class FakeEKTClient:
    def __init__(self, products, details):
        self.products = products
        self.details = details

    async def get_products(self, page=1):
        return {
            "data": self.products if page == 1 else [],
            "meta": {"last_page": 1},
        }

    async def get_product_detail(self, product_id):
        return {"data": self.details[product_id]}


def run_immediate(coroutine):
    """Drive fake-I/O coroutines without opening sockets in restricted CI."""
    try:
        coroutine.send(None)
    except StopIteration as exc:
        return exc.value
    raise AssertionError("The test coroutine unexpectedly performed real I/O")


class CatalogTests(unittest.TestCase):
    def test_supplier_article_has_priority_and_stock_is_checked(self):
        summary = {
            "id": 515291,
            "name": "027228 АВ DRX250 MT 3ф 160А 18ka Legrand (1)",
            "article": "200300285_",
            "properties": {"ARTIKULPOSTAVSHCHIKA": "027228"},
        }
        detail = {
            **summary,
            "price": 64920,
            "quantity": 23,
            "stores": [
                {"name": "Алматы", "quantity": 5},
                {"name": "Нур-Султан", "quantity": 0},
            ],
            "properties": {
                "ARTIKULPOSTAVSHCHIKA": "027228",
                "NOMINALNYY_TOK": "250 А",
            },
        }
        service = CatalogService(FakeEKTClient([summary], {515291: detail}))
        item = SpecificationItem(
            raw_text="027228 | Legrand DRX250 | 30",
            search_query="027228 Legrand DRX250 160A",
            article="027228",
            quantity=30,
            unit="шт",
        )

        result = run_immediate(service.match_item(item))

        self.assertEqual("matched", result.match.status)
        self.assertEqual("027228", result.match.product.supplier_article)
        self.assertEqual(23, result.available_quantity)
        self.assertFalse(result.enough_stock)
        self.assertEqual(["Алматы"], [store.name for store in result.match.product.stores])
        self.assertEqual("nominal_current", result.match.product.data_warnings[0].field)

    def test_duplicate_exact_articles_are_ambiguous(self):
        products = [
            {
                "id": 1,
                "name": "Product one",
                "article": "internal-1",
                "properties": {"ARTIKULPOSTAVSHCHIKA": "027228"},
            },
            {
                "id": 2,
                "name": "Product two",
                "article": "internal-2",
                "properties": {"ARTIKULPOSTAVSHCHIKA": "027228"},
            },
        ]
        service = CatalogService(
            FakeEKTClient(products, {1: products[0], 2: products[1]})
        )
        item = SpecificationItem(
            raw_text="027228",
            search_query="027228",
            article="027228",
        )

        result = run_immediate(service.match_item(item))

        self.assertEqual("ambiguous", result.match.status)
        self.assertEqual(2, len(result.match.candidates))


if __name__ == "__main__":
    unittest.main()
